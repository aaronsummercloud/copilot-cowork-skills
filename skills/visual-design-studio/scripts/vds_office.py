"""Office builders for visual-design-studio: Word (.docx) and Excel (.xlsx). Requires python-docx and openpyxl.

=== WORD ===
Build a designed Word document from Markdown and the design tokens.

Usage: python3 scripts/vds.py docx source.md out.docx [--theme theme.json]

Requires python-docx (preinstalled in Copilot Cowork). The theme maps the DESIGN.md tokens to
Office-safe values. Put the theme keys in the Markdown front matter (preferred), or pass a JSON file
with --theme. Keys: font_display, font_body (Office fonts), ink, ink_muted, accent, accent_ink,
surface_raised, line (hex without #), size_body, size_h1, size_h2, size_h3 (pt), margin_mm.
Missing keys fall back to a neutral default.

Markdown subset:
  ---                      front matter at the top: title, subtitle, org, date, author, plus any
                           theme keys (key: value lines)
  # / ## / ###             headings (H1 starts a new page, except the first)
  plain lines              paragraphs; **bold**, *italic*
  - item / 1. item         bullet and numbered lists
  > text                   pull quote; a following "> -- Name" line becomes the attribution
  > [!NOTE] Title          callout box; following "> " lines are its body
  !stat 18.4% | caption    key-figure block (one big number plus a caption)
  | a | b |                table; the first row is the header; a |---| separator line is ignored
  \\pagebreak               explicit page break
  ![alt](chart.png){width=150}   image (path relative to the .md; width in mm, default full text width),
                           captioned with its alt text; make charts with `vds.py chart data.json chart.png`
  [TOC]                    table of contents (Word fills it in when the file is opened)

=== EXCEL ===
Build a designed, live Excel dashboard workbook from a JSON spec.

Usage: python3 scripts/vds.py xlsx spec.json out.xlsx

Requires openpyxl (preinstalled in Copilot Cowork). See examples/dashboard-harbour.json for the
spec format and reference/layout-patterns.md (Dashboards) for the design rules. The output has:
  - Data sheets: one formatted Excel Table per spec table (typed number formats, frozen header,
    filters, optional data bars and color scales). This is where users edit numbers.
  - A Dashboard sheet on a 12-column grid: title block, KPI tiles, chart/table/note panels.
    Every KPI and chart is a live formula or reference to the data, so editing data updates it.
Mixed units in one table: give the table "row_formats": one number format per row (or null), which overrides
the value cells of columns that have no "format" of their own (e.g. euros in one row and hours in another,
while a "Change" column keeps its percentage format).
Formula placeholders: {Table.Column} -> the column's data range; {Table.Column[last]} -> its last
cell; {Table.Column[-2]} -> second-to-last; {Table.Column[0]} -> first.
Excel recalculates on open (fullCalcOnLoad). For QA in Cowork, render with
scripts/vds.py render out.xlsx, which also flags formula errors.
"""
import json
import os
import re
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.layout import Layout, ManualLayout
    from openpyxl.chart.legend import Legend
    from openpyxl.chart.marker import DataPoint, Marker
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.chart.text import RichText
    from openpyxl.chart.axis import ChartLines
    from openpyxl.drawing.line import LineProperties
    from openpyxl.drawing.text import CharacterProperties, Font as DrawFont, Paragraph, ParagraphProperties
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.worksheet.table import Table, TableStyleInfo
    _OFFICE_MISSING = None
except ImportError as _exc:  # deck-check / lint paths don't need these
    _OFFICE_MISSING = _exc


# ============================================================================
# Word: Markdown subset -> designed .docx (python-docx)
# ============================================================================
DOCX_THEME = {
    "font_display": "Calibri", "font_body": "Cambria",
    "ink": "15211B", "ink_muted": "4C5D54", "accent": "1F6B47", "accent_ink": "FFFFFF",
    "surface_raised": "EEF4EF", "line": "CFDAD3",
    "size_body": 10.5, "size_h1": 24, "size_h2": 16, "size_h3": 12, "margin_mm": 22,
}


def rgb(hexv):
    return RGBColor.from_string(hexv.lstrip("#").upper())


TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]


def tcpr_insert(tc_pr, el):
    """Insert a child into <w:tcPr> in schema order (Word rejects out-of-order cell properties)."""
    name = el.tag.split("}")[1]
    rank = TCPR_ORDER.index(name)
    for old in tc_pr.findall(qn(f"w:{name}")):
        tc_pr.remove(old)
    for i, child in enumerate(tc_pr):
        child_name = child.tag.split("}")[1]
        if child_name in TCPR_ORDER and TCPR_ORDER.index(child_name) > rank:
            tc_pr.insert(i, el)
            return
    tc_pr.append(el)


def shade(cell, hexv):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexv.lstrip("#").upper())
    tcpr_insert(tc_pr, shd)


def cell_margins(cell, mm):
    """Inner padding for a table cell, in millimetres (top, right, bottom, left)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    top, right, bottom, left = mm
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):  # schema order
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(value * 56.7)))  # 1 mm = 56.7 twips
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr_insert(tc_pr, mar)


def cell_borders(cell, **edges):
    """edges: top/bottom/left/right = (size_eighths_pt, hex) or None for none."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:color"), spec[1].lstrip("#").upper())
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcpr_insert(tc_pr, borders)


def para_border_left(paragraph, hexv, size=18, space=12):
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", str(space)), ("color", hexv.lstrip("#").upper())):
        left.set(qn(f"w:{k}"), v)
    bdr.append(left)
    p_pr.append(bdr)


def add_field(run, instr):
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if attr:
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)


def add_inline(paragraph, text, theme, size=None, color=None, font=None, bold=None, italic=None):
    """Write text with **bold** and *italic* markup as runs."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        run = paragraph.add_run(part.strip("*") if part.startswith("*") else part)
        run.bold = True if part.startswith("**") else bold
        run.italic = True if (part.startswith("*") and not part.startswith("**")) else italic
        run.font.name = font or theme["font_body"]
        run.font.size = Pt(size or theme["size_body"])
        run.font.color.rgb = rgb(color or theme["ink"])
    return paragraph


def setup_styles(doc, theme):
    normal = doc.styles["Normal"]
    normal.font.name = theme["font_body"]
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), theme["font_body"])
    normal.font.size = Pt(theme["size_body"])
    normal.font.color.rgb = rgb(theme["ink"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    for level, size in ((1, theme["size_h1"]), (2, theme["size_h2"]), (3, theme["size_h3"])):
        style = doc.styles[f"Heading {level}"]
        style.font.name = theme["font_display"]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(theme["ink"] if level != 3 else theme["accent"])
        style.paragraph_format.space_before = Pt(18 if level == 1 else 14)
        style.paragraph_format.space_after = Pt(8 if level == 1 else 4)
        style.paragraph_format.keep_with_next = True
        r_fonts = style.element.rPr.rFonts
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            r_fonts.set(qn(attr), theme["font_display"])
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if r_fonts.get(qn(attr)) is not None:
                del r_fonts.attrib[qn(attr)]
    for name in ("List Bullet", "List Number"):
        doc.styles[name].font.name = theme["font_body"]
        doc.styles[name].font.size = Pt(theme["size_body"])


def cover(doc, meta, theme):
    """Full-width accent band with the title, then the subtitle and metadata."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, theme["accent"])
    cell_borders(cell)
    cell_margins(cell, (14, 14, 14, 14))
    cell.text = ""
    top = cell.paragraphs[0]
    add_inline(top, meta.get("org", ""), theme, size=9, color=theme["accent_ink"], font=theme["font_display"], bold=True)
    top.paragraph_format.space_after = Pt(120)
    title = cell.add_paragraph()
    add_inline(title, meta.get("title", "Untitled"), theme, size=40, color=theme["accent_ink"],
               font=theme["font_display"], bold=True)
    title.paragraph_format.line_spacing = 0.95
    title.paragraph_format.space_after = Pt(10)
    if meta.get("subtitle"):
        sub = cell.add_paragraph()
        add_inline(sub, meta["subtitle"], theme, size=14, color=theme["accent_ink"])
        sub.paragraph_format.space_after = Pt(28)
    row = table.rows[0]
    row.height = Mm(200)
    info = doc.add_paragraph()
    info.paragraph_format.space_before = Pt(14)
    add_inline(info, "   ".join(x for x in (meta.get("author"), meta.get("date")) if x), theme,
               size=10, color=theme["ink_muted"], font=theme["font_display"])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def stat(doc, value, caption, theme):
    table = doc.add_table(rows=1, cols=2)
    a, b = table.cell(0, 0), table.cell(0, 1)
    for c in (a, b):
        cell_borders(c)
    a.width, b.width = Mm(55), Mm(110)
    add_inline(a.paragraphs[0], value, theme, size=34, color=theme["accent"], font=theme["font_display"], bold=True)
    add_inline(b.paragraphs[0], caption, theme, size=12, color=theme["ink_muted"])
    b.vertical_alignment = 1
    doc.add_paragraph()


def callout(doc, title, body_lines, theme):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, theme["surface_raised"])
    cell_borders(cell, left=(24, theme["accent"]))
    cell_margins(cell, (3, 4, 3, 5))
    head = cell.paragraphs[0]
    add_inline(head, title, theme, size=theme["size_body"] + 0.5, font=theme["font_display"], bold=True,
               color=theme["accent"])
    for line in body_lines:
        add_inline(cell.add_paragraph(), line, theme, size=theme["size_body"] - 0.5)
    doc.add_paragraph()


def pull_quote(doc, text, attribution, theme):
    p = doc.add_paragraph()
    para_border_left(p, theme["accent"])
    p.paragraph_format.left_indent = Mm(4)
    p.paragraph_format.space_before = Pt(10)
    add_inline(p, text, theme, size=15, italic=True)
    if attribution:
        a = doc.add_paragraph()
        para_border_left(a, theme["accent"])
        a.paragraph_format.left_indent = Mm(4)
        a.paragraph_format.space_after = Pt(12)
        add_inline(a, attribution, theme, size=8.5, color=theme["ink_muted"], font=theme["font_display"])


def table_block(doc, rows, theme):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=len(rows), cols=len(header))
    numeric = [all(re.fullmatch(r"[−\-+£$€]?[\d.,]+%?\s*\w{0,3}", r[i].strip() or "0") for r in body)
               for i in range(len(header))]
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell_borders(cell, bottom=(12 if r == 0 else 4, theme["ink"] if r == 0 else theme["line"]))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if numeric[c] and c > 0 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value.strip(), theme, size=theme["size_body"] - 1, font=theme["font_display"],
                       bold=(r == 0) or None, color=theme["ink_muted"] if r == 0 else theme["ink"])
    doc.add_paragraph()


def image(doc, alt, path, width_mm, theme):
    """A figure: the picture at the requested width, alt text for screen readers, alt text as caption."""
    usable = 210 - 2 * theme["margin_mm"]
    doc.add_picture(path, width=Mm(min(width_mm or usable, usable)))
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for doc_pr in pic._p.iter(qn("wp:docPr")):
        doc_pr.set("descr", alt)
    if alt:
        caption = doc.add_paragraph()
        add_inline(caption, alt, theme, size=8.5, color=theme["ink_muted"], font=theme["font_display"])
        caption.paragraph_format.space_after = Pt(10)


def toc(doc, theme, headings):
    """A TOC field pre-filled with the headings, so viewers that don't update fields (LibreOffice, previews)
    still show the contents; settings.updateFields makes Word rebuild it with page numbers on open."""
    p = doc.add_paragraph()
    add_inline(p, "Contents", theme, size=theme["size_h2"], font=theme["font_display"], bold=True)
    entries = headings or [(1, "")]
    paras = []
    for level, text in entries:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Mm(6 * (level - 1))
        para.paragraph_format.space_after = Pt(3)
        paras.append((para, text, level))
    first = paras[0][0]

    def fld(paragraph, kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(el)

    fld(first, "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    first.add_run()._r.append(instr)
    fld(first, "separate")
    for para, text, level in paras:
        add_inline(para, text, theme, size=theme["size_body"] + (1 if level == 1 else 0),
                   bold=(level == 1) or None, color=theme["ink"] if level == 1 else theme["ink_muted"])
    fld(paras[-1][0], "end")
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        # schema order: updateFields comes before these elements (Word rejects it elsewhere)
        after = ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr",
                 "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                 "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType",
                 "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")
        anchor = next((child for child in settings if child.tag.split("}")[-1] in after), None)
        if anchor is not None:
            anchor.addprevious(upd)
        else:
            settings.append(upd)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def footer(section, title, theme):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    add_inline(p, title + "    ", theme, size=8, color=theme["ink_muted"], font=theme["font_display"])
    run = p.add_run()
    run.font.size = Pt(8)
    run.font.name = theme["font_display"]
    run.font.color.rgb = rgb(theme["ink_muted"])
    add_field(run, "PAGE")


def parse_front_matter(lines):
    meta = {}
    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1) if "---" in lines[1:] else 0
        for line in lines[1:end]:
            if ":" in line:
                key, value = line.split(":", 1)
                meta[key.strip().lower()] = value.strip()
        lines = lines[end + 1:]
    return meta, lines


def build_docx(src, out, theme):
    lines = [l.rstrip("\n") for l in open(src, encoding="utf-8")]
    meta, lines = parse_front_matter([l.rstrip() for l in lines])
    toc_headings = [(len(m.group(1)), m.group(2).strip()) for m in
                    (re.match(r"(#{1,2})\s+(.*)", l.strip()) for l in lines) if m]
    theme = dict(theme)
    for key in DOCX_THEME:
        if key in meta:
            value = meta[key].lstrip("#")
            theme[key] = float(value) if key.startswith(("size_", "margin")) else value
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(theme["margin_mm"]))
    setup_styles(doc, theme)
    doc.core_properties.title = meta.get("title", "")
    doc.core_properties.author = meta.get("author", "")
    if meta.get("title"):
        cover(doc, meta, theme)
        section.different_first_page_header_footer = True  # no running footer on the cover
    footer(section, meta.get("title", ""), theme)

    i, first_h1, para_buf = 0, True, []

    def flush():
        if para_buf:
            add_inline(doc.add_paragraph(), " ".join(para_buf), theme)
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush()
        elif re.match(r"!\[[^\]]*\]\([^)]+\)", s):
            flush()
            m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)(?:\{width=(\d+(?:\.\d+)?)\})?", s)
            image(doc, m.group(1), os.path.join(os.path.dirname(os.path.abspath(src)), m.group(2)),
                  float(m.group(3)) if m.group(3) else None, theme)
        elif s == "[TOC]":
            flush()
            toc(doc, theme, toc_headings)
        elif s == "\\pagebreak":
            flush()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            if level == 1 and not first_h1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_h1 = first_h1 and level != 1
            doc.add_heading(s[level:].strip(), level=min(level, 3))
        elif s.startswith("!stat "):
            flush()
            value, _, caption = s[6:].partition("|")
            stat(doc, value.strip(), caption.strip(), theme)
        elif s.startswith("> [!"):
            flush()
            title = re.sub(r"^> \[![A-Z]+\]\s*", "", s) or "Note"
            body = []
            while i + 1 < len(lines) and lines[i + 1].strip().startswith(">"):
                i += 1
                body.append(lines[i].strip()[1:].strip())
            callout(doc, title, [b for b in body if b], theme)
        elif s.startswith(">"):
            flush()
            text, attribution = s[1:].strip(), None
            if i + 1 < len(lines) and re.match(r">\s*(--|-|—)\s*", lines[i + 1].strip()):
                i += 1
                attribution = re.sub(r">\s*(--|-|—)\s*", "", lines[i].strip())
            pull_quote(doc, text, attribution, theme)
        elif re.match(r"[-*] ", s):
            flush()
            add_inline(doc.add_paragraph(style="List Bullet"), s[2:], theme)
        elif re.match(r"\d+[.)] ", s):
            flush()
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)] ", "", s), theme)
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            table_block(doc, rows, theme)
        else:
            para_buf.append(s)
        i += 1
    flush()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # python-docx's default template omits this required attribute
    doc.save(out)


def main_docx():
    if _OFFICE_MISSING:
        sys.exit(f"python-docx and openpyxl are required ({_OFFICE_MISSING})")
    argv = sys.argv[1:]
    theme_path = None
    if "--theme" in argv:
        k = argv.index("--theme")
        theme_path = argv[k + 1]
        del argv[k:k + 2]
    args = [a for a in argv if not a.startswith("--")]
    if len(args) != 2:
        print(__doc__)
        sys.exit(2)
    theme = dict(DOCX_THEME)
    if theme_path:
        with open(theme_path, encoding="utf-8") as fh:
            theme.update({k: v for k, v in json.load(fh).items() if not k.startswith("_")})
    build_docx(args[0], args[1], theme)
    print(f"Wrote {args[1]}")



# ============================================================================
# Excel: JSON spec -> live dashboard workbook (openpyxl)
# ============================================================================
XLSX_THEME = {
    "font_display": "Segoe UI Semibold", "font_body": "Segoe UI",
    "ink": "1B2430", "ink_muted": "5A6573", "surface": "FFFFFF", "surface_raised": "EEF1F5",
    "line": "D5DBE3", "accent": "2F5D8A", "accent_ink": "FFFFFF", "highlight": "E0892B",
    "good": "2E7D4F", "bad": "B23A3A", "series": ["2F5D8A", "9DB4CC", "E0892B", "5A6573", "7FA27A", "C9A227"],
}
GRID_COLS, COL_W, GUTTER_W, ROW_PT = 12, 11.0, 2.5, 20
BANNED = ["XLOOKUP", "XMATCH", "SORT(", "SORTBY", "FILTER(", "UNIQUE(", "SEQUENCE(", "LET(", "LAMBDA("]
XLFN = ["TEXTJOIN", "CONCAT", "IFS", "SWITCH", "MAXIFS", "MINIFS"]
warnings = []


# ---------------------------------------------------------------- helpers

def side(color, style="thin"):
    return Side(style=style, color=color)


def text_props(size_pt, color, face, bold=False):
    cp = CharacterProperties(sz=int(size_pt * 100), b=bold, solidFill=color, latin=DrawFont(typeface=face))
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def col_px(width_chars):
    return int(width_chars * 7 + 5)


class Tables:
    """Where each spec table landed, for formula placeholders and chart references."""

    def __init__(self):
        self.by_name = {}

    def add(self, name, ws, columns, first_row, n_rows):
        self.by_name[name] = {"ws": ws, "cols": {c["name"]: i + 1 for i, c in enumerate(columns)},
                              "first": first_row, "last": first_row + n_rows - 1, "spec": columns}

    def col_ref(self, table, column):
        t = self.by_name[table]
        return t["ws"], t["cols"][column], t["first"], t["last"]

    def expand(self, formula):
        def repl(m):
            table, column, idx = m.group(1), m.group(2), m.group(3)
            ws, c, first, last = self.col_ref(table, column)
            sheet = f"'{ws.title}'" if " " in ws.title else ws.title
            letter = get_column_letter(c)
            if idx is None:
                return f"{sheet}!${letter}${first}:${letter}${last}"
            row = last if idx == "last" else (last + 1 + int(idx) if int(idx) < 0 else first + int(idx))
            return f"{sheet}!${letter}${row}"
        out = re.sub(r"\{([\w ]+?)\.([\w .%()/£$€-]+?)(?:\[(last|-?\d+)\])?\}", repl, formula)
        upper = out.upper()
        for fn in BANNED:
            if fn in upper:
                raise ValueError(f"formula uses {fn.rstrip('(')}: LibreOffice can't evaluate it; "
                                 "use INDEX/MATCH or precompute in Python")
        for fn in XLFN:
            out = re.sub(rf"(?<![\w.]){fn}\(", f"_xlfn.{fn}(", out)
        return out


# ---------------------------------------------------------------- data sheets

def build_tables(wb, spec, theme, tables):
    hdr_font = Font(name=theme["font_display"], size=10, color=theme["ink"], bold=True)
    body_font = Font(name=theme["font_body"], size=10, color=theme["ink"])
    sheets = {}
    for t in spec["tables"]:
        ws = sheets.get(t.get("sheet", "Data"))
        if ws is None:
            ws = wb.create_sheet(t.get("sheet", "Data"))
            ws.sheet_view.showGridLines = False
            sheets[ws.title] = ws
            ws._next_col = 1
        start_col = ws._next_col
        cols, rows = t["columns"], t["rows"]
        for j, c in enumerate(cols):
            cell = ws.cell(row=1, column=start_col + j, value=c["name"])
            cell.font, cell.border = hdr_font, Border(bottom=side(theme["ink"], "medium"))
            cell.alignment = Alignment(horizontal="right" if c.get("format") and j else "left", vertical="center")
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = ws.cell(row=2 + i, column=start_col + j, value=tables.expand(value) if isinstance(value, str) and value.startswith("=") else value)
                cell.font = body_font
                cell.border = Border(bottom=side(theme["line"], "hair"))
                row_fmt = (t.get("row_formats") or [None] * len(rows))[i]
                if row_fmt and j > 0 and not cols[j].get("format"):  # columns with their own format keep it
                    cell.number_format = row_fmt
                elif cols[j].get("format"):
                    cell.number_format = cols[j]["format"]
        for j, c in enumerate(cols):
            longest = max([len(str(c["name"]))] + [len(str(r[j])) for r in rows])
            ws.column_dimensions[get_column_letter(start_col + j)].width = min(max(longest * 1.15 + 3, 10), 40)
        ref = f"{get_column_letter(start_col)}1:{get_column_letter(start_col + len(cols) - 1)}{len(rows) + 1}"
        table = Table(displayName=re.sub(r"\W", "", t["name"]), ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(table)
        tables.add(t["name"], ws, cols, 2, len(rows))
        for j, c in enumerate(cols):
            rng = f"{get_column_letter(start_col + j)}2:{get_column_letter(start_col + j)}{len(rows) + 1}"
            if c.get("bar"):
                ws.conditional_formatting.add(rng, DataBarRule(start_type="min", end_type="max", color=theme["accent"]))
            if c.get("scale"):
                ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color=theme["surface_raised"],
                                                                  end_type="max", end_color=theme["accent"]))
        ws.freeze_panes = "A2"
        ws._next_col = start_col + len(cols) + 1
    return sheets


# ---------------------------------------------------------------- dashboard

def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def block(ws, top, left, bottom, right, color, gap_color=None):
    for r in range(top, bottom + 1):
        for c in range(left, right + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill(color)
            if gap_color:
                cell.border = Border(right=side(gap_color, "thick") if c == right else None)


C_NS = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"


def suppress_auto_title(chart):
    """openpyxl never writes <c:autoTitleDeleted>, so Excel titles single-series charts with the series
    name. The panel title carries the message, so delete the automatic one."""
    original = chart._write

    def _write():
        tree = original()
        # openpyxl writes child tags unqualified (the namespace sits on the root), so match both forms.
        ns = C_NS if tree.find(C_NS + "chart") is not None else ""
        container = tree.find(ns + "chart")
        if container is not None and container.find(ns + "autoTitleDeleted") is None:
            index = 1 if container.find(ns + "title") is not None else 0
            # makeelement works for both lxml and xml.etree trees (openpyxl uses whichever is installed)
            container.insert(index, container.makeelement(ns + "autoTitleDeleted", {"val": "1"}))
        return tree
    chart._write = _write


def build_chart(spec, theme, tables, width_cm, height_cm):
    kind = spec["type"]
    t = tables.by_name[spec["table"]]
    ws = t["ws"]
    x_col = t["cols"][spec["x"]]
    ys = spec["y"] if isinstance(spec["y"], list) else [spec["y"]]
    if len(ys) > 6:
        warnings.append(f"chart '{spec.get('title', '')}': more than 6 series")
    if kind in ("pie", "doughnut"):
        chart = PieChart() if kind == "pie" else DoughnutChart(holeSize=62)
        if t["last"] - t["first"] + 1 > 5:
            warnings.append("pie/doughnut with more than 5 slices: use a bar chart")
    elif kind == "line":
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = "bar" if kind.startswith("bar") or kind.startswith("stacked-bar") else "col"
        chart.grouping = "stacked" if kind.startswith("stacked") else "clustered"
        if kind.startswith("stacked"):
            chart.overlap = 100
        chart.gapWidth = spec.get("gap_width", 60)
    for y in ys:
        c = t["cols"][y]
        chart.add_data(Reference(ws, min_col=c, min_row=t["first"] - 1, max_row=t["last"]), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=x_col, min_row=t["first"], max_row=t["last"]))
    colors = spec.get("colors") or theme["series"]
    chart.title = None
    chart.style = 2
    chart.roundedCorners = False
    chart.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    suppress_auto_title(chart)
    chart.width, chart.height = width_cm, height_cm
    ink, muted, face = theme["ink"], theme["ink_muted"], theme["font_body"]
    if kind in ("pie", "doughnut"):
        series = chart.series[0]
        for i in range(t["last"] - t["first"] + 1):
            pt = DataPoint(idx=i)
            pt.graphicalProperties.solidFill = colors[i % len(colors)]
            pt.graphicalProperties.line.solidFill = theme["surface"]
            pt.graphicalProperties.line.width = 19050
            series.dPt.append(pt)
        chart.legend = Legend()
        chart.legend.position = "r"
        chart.legend.txPr = text_props(9, muted, face)
    else:
        for i, series in enumerate(chart.series):
            color = colors[i % len(colors)]
            if kind == "line":
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 28575
                series.smooth = False
                series.marker = Marker(symbol="circle", size=6)
                series.marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color))
            else:
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.noFill = True
        hl = spec.get("highlight")
        if hl is not None and kind != "line":
            n = t["last"] - t["first"] + 1
            idx = n - 1 if hl == "last" else int(hl)
            pt = DataPoint(idx=idx)
            pt.graphicalProperties.solidFill = theme["highlight"]
            pt.graphicalProperties.line.noFill = True
            chart.series[0].dPt.append(pt)
        y, x = chart.y_axis, chart.x_axis
        y.number_format = spec.get("format", "General")
        y.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=theme["line"], w=6350)))
        y.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        x.spPr = GraphicalProperties(ln=LineProperties(solidFill=theme["line"], w=9525))
        x.majorGridlines = None
        for axis in (x, y):
            axis.txPr = text_props(9, muted, face)
            axis.delete = False
        y.delete = bool(spec.get("labels")) and kind != "line"
        if len(ys) > 1:
            chart.legend = Legend()
            chart.legend.position = "t"
            chart.legend.txPr = text_props(9, muted, face)
        else:
            chart.legend = None
    if spec.get("labels"):
        chart.dataLabels = DataLabelList(showVal=True)
        chart.dataLabels.numFmt = spec.get("format", "General")
        chart.dataLabels.txPr = text_props(9, ink, face)
        if kind in ("pie", "doughnut"):
            chart.dataLabels.showLeaderLines = False
    chart.layout = Layout(manualLayout=ManualLayout(x=0.02, y=0.04, w=0.96, h=0.9)) if kind not in ("pie", "doughnut") else None
    return chart


def build_dashboard(wb, spec, theme, tables):
    d = spec["dashboard"]
    ws = wb.create_sheet(d.get("sheet", "Dashboard"), 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    ws.sheet_properties.tabColor = theme["accent"]
    ws.column_dimensions["A"].width = GUTTER_W
    for c in range(2, GRID_COLS + 2):
        ws.column_dimensions[get_column_letter(c)].width = COL_W
    ws.column_dimensions[get_column_letter(GRID_COLS + 2)].width = GUTTER_W
    for r in range(1, 200):
        ws.row_dimensions[r].height = ROW_PT
    disp, body = theme["font_display"], theme["font_body"]
    last_col = get_column_letter(GRID_COLS + 1)

    # Title block: claim title, subtitle, meta right-aligned
    ws.row_dimensions[2].height = 36
    ws.merge_cells(f"B2:{get_column_letter(GRID_COLS - 2)}2")
    ws["B2"] = spec["title"]
    ws["B2"].font = Font(name=disp, size=22, bold=True, color=theme["ink"])
    ws["B2"].alignment = Alignment(vertical="bottom")
    ws.merge_cells(f"B3:{get_column_letter(GRID_COLS - 2)}3")
    ws["B3"] = spec.get("subtitle", "")
    ws["B3"].font = Font(name=body, size=11, color=theme["ink_muted"])
    meta_col = get_column_letter(GRID_COLS - 1)
    ws.merge_cells(f"{meta_col}3:{last_col}3")
    ws[f"{meta_col}3"] = f"Updated {spec['updated']}" if spec.get("updated") else ""
    ws[f"{meta_col}3"].font = Font(name=body, size=9, color=theme["ink_muted"])
    ws[f"{meta_col}3"].alignment = Alignment(horizontal="right")
    for c in range(2, GRID_COLS + 2):
        ws.cell(row=4, column=c).border = Border(bottom=side(theme["line"]))

    # KPI tiles: one band, equal spans, the lead tile may use the accent
    row = 6
    kpis = d.get("kpis", [])
    if len(kpis) > 6:
        warnings.append("more than 6 KPIs: pick the ones that drive decisions")
    if kpis:
        span = GRID_COLS // len(kpis)
        for i, k in enumerate(kpis):
            left = 2 + i * span
            right = left + span - 1 if i < len(kpis) - 1 else GRID_COLS + 1
            lead = k.get("lead", False)
            bg, fg = (theme["accent"], theme["accent_ink"]) if lead else (theme["surface_raised"], theme["ink"])
            muted = theme["accent_ink"] if lead else theme["ink_muted"]
            block(ws, row, left, row + 3, right, bg, gap_color=theme["surface"] if i < len(kpis) - 1 else None)
            L, R = get_column_letter(left), get_column_letter(right)
            for r_off, height in ((0, 22), (1, 40), (2, 20), (3, 10)):
                ws.row_dimensions[row + r_off].height = height
            ws.merge_cells(f"{L}{row}:{R}{row}")
            ws[f"{L}{row}"] = k["label"]
            ws[f"{L}{row}"].font = Font(name=body, size=10, color=muted)
            ws[f"{L}{row}"].alignment = Alignment(horizontal="left", indent=1, vertical="bottom")
            ws.merge_cells(f"{L}{row + 3}:{R}{row + 3}")
            ws.merge_cells(f"{L}{row + 1}:{R}{row + 1}")
            vcell = ws[f"{L}{row + 1}"]
            vcell.value = tables.expand(k["value"]) if str(k["value"]).startswith("=") else k["value"]
            vcell.number_format = k.get("format", "General")
            vcell.font = Font(name=disp, size=26, bold=True, color=fg)
            vcell.alignment = Alignment(horizontal="left", indent=1, vertical="center")
            if k.get("delta"):
                ws.merge_cells(f"{L}{row + 2}:{R}{row + 2}")
                dcell = ws[f"{L}{row + 2}"]
                dcell.value = tables.expand(k["delta"])
                dcell.number_format = k.get("delta_format", '+0.0%;-0.0%;0.0%')
                dcell.font = Font(name=body, size=10, bold=True, color=muted)
                dcell.alignment = Alignment(horizontal="left", indent=1, vertical="top")
                if not lead:
                    good, bad = theme["good"], theme["bad"]
                    up_is_good = k.get("good", "up") == "up"
                    ws.conditional_formatting.add(dcell.coordinate, CellIsRule(
                        operator="greaterThan", formula=["0"], font=Font(color=good if up_is_good else bad, bold=True)))
                    ws.conditional_formatting.add(dcell.coordinate, CellIsRule(
                        operator="lessThan", formula=["0"], font=Font(color=bad if up_is_good else good, bold=True)))
        row += 5

    # Panels: flow layout on the 12-column grid
    col, band_bottom = 0, row
    col_cm = col_px(COL_W) / 37.795
    row_cm = ROW_PT * 4 / 3 / 37.795
    for p in d.get("panels", []):
        w, h = p.get("width", 6), p.get("height", 14)
        if col + w > GRID_COLS:
            col, row = 0, band_bottom + 1
        left = 2 + col
        L, R = get_column_letter(left), get_column_letter(left + w - 1)
        ws.merge_cells(f"{L}{row}:{R}{row}")
        ws[f"{L}{row}"] = p["title"]
        ws[f"{L}{row}"].font = Font(name=disp, size=12, bold=True, color=theme["ink"])
        ws[f"{L}{row}"].alignment = Alignment(vertical="bottom")
        if p.get("caption"):
            ws.merge_cells(f"{L}{row + 1}:{R}{row + 1}")
            ws[f"{L}{row + 1}"] = p["caption"]
            ws[f"{L}{row + 1}"].font = Font(name=body, size=9, color=theme["ink_muted"])
        top = row + 2
        if "chart" in p:
            # Anchor cell-to-cell: apps convert column widths to pixels differently, so a chart sized in
            # cm overruns its span in some viewers. A two-cell anchor covers exactly the panel's cells.
            chart = build_chart(p["chart"], theme, tables, w * col_cm, h * row_cm)
            chart.anchor = TwoCellAnchor(_from=AnchorMarker(col=left - 1, row=top - 1),
                                         to=AnchorMarker(col=left - 1 + w, row=top - 1 + h))
            ws.add_chart(chart)
        elif "table" in p:
            tp = p["table"]
            ws_t, _, first, last = tables.col_ref(tp["table"], tp["columns"][0])
            n = min(tp.get("rows", last - first + 1), last - first + 1)
            spans = tp.get("spans") or [w // len(tp["columns"])] * len(tp["columns"])
            c0 = left
            for j, cname in enumerate(tp["columns"]):
                cl = get_column_letter(c0)
                cr = get_column_letter(c0 + spans[j] - 1)
                # Numbers go in the LAST cell of their span (right-aligned); text in the first (left-aligned),
                # overflowing into the empty cells. No merges: merged cells clip in some viewers.
                num = bool(tables.by_name[tp["table"]]["spec"][tables.by_name[tp["table"]]["cols"][cname] - 1].get("format"))
                cl = cr if (num and j) else cl
                hdr = ws[f"{cl}{top}"]
                hdr.value = cname
                hdr.font = Font(name=disp, size=9, bold=True, color=theme["ink_muted"])
                hdr.alignment = Alignment(horizontal="right" if num and j else "left", indent=0 if num and j else 1)
                for c in range(c0, c0 + spans[j]):
                    ws.cell(row=top, column=c).border = Border(bottom=side(theme["ink"], "medium"))
                _, src_col, src_first, _ = tables.col_ref(tp["table"], cname)
                sheet = f"'{ws_t.title}'" if " " in ws_t.title else ws_t.title
                for i in range(n):
                    r = top + 1 + i
                    cell = ws[f"{cl}{r}"]
                    cell.value = f"={sheet}!{get_column_letter(src_col)}{src_first + i}"
                    fmt = ws_t.cell(row=src_first + i, column=src_col).number_format  # row formats carry over
                    if fmt and fmt != "General":
                        cell.number_format = fmt
                    cell.font = Font(name=body, size=10, color=theme["ink"])
                    cell.alignment = Alignment(horizontal="right" if num and j else "left", indent=0 if num and j else 1)
                    for c in range(c0, c0 + spans[j]):
                        ws.cell(row=r, column=c).border = Border(bottom=side(theme["line"], "hair"))
                if cname == tp.get("bar"):
                    ws.conditional_formatting.add(f"{cl}{top + 1}:{cl}{top + n}", DataBarRule(
                        start_type="num", start_value=0, end_type="max", color=theme["series"][1]))
                c0 += spans[j]
            h = max(h, n + 1)
        elif "text" in p:
            ws.merge_cells(f"{L}{top}:{R}{top + h - 1}")
            cell = ws[f"{L}{top}"]
            cell.value = p["text"]
            cell.font = Font(name=body, size=10, color=theme["ink"])
            cell.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            block(ws, top, left, top + h - 1, left + w - 1, theme["surface_raised"])
        band_bottom = max(band_bottom, top + h)
        col += w

    # Footer: source line
    foot = band_bottom + 1
    if spec.get("source"):
        ws.merge_cells(f"B{foot}:{last_col}{foot}")
        ws[f"B{foot}"] = f"Source: {spec['source']}"
        ws[f"B{foot}"].font = Font(name=body, size=9, color=theme["ink_muted"])
    # Print: landscape, one page wide
    ws.print_area = f"A1:{get_column_letter(GRID_COLS + 2)}{foot + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    return ws


def build_xlsx(spec, out):
    theme = dict(XLSX_THEME)
    theme.update(spec.get("theme", {}))
    theme = {k: (v.lstrip("#").upper() if isinstance(v, str) and re.fullmatch(r"#?[0-9a-fA-F]{6}", v) else v)
             for k, v in theme.items()}
    theme["series"] = [s.lstrip("#").upper() for s in theme["series"]]
    wb = Workbook()
    wb.remove(wb.active)
    tables = Tables()
    build_tables(wb, spec, theme, tables)
    dash = build_dashboard(wb, spec, theme, tables)
    wb.active = wb.index(dash)
    wb.properties.title = spec["title"]
    wb.properties.creator = spec.get("author", "")
    wb.calculation.fullCalcOnLoad = True  # openpyxl stores no cached values; Excel computes on open
    wb.save(out)


def main_xlsx():
    if _OFFICE_MISSING:
        sys.exit(f"python-docx and openpyxl are required ({_OFFICE_MISSING})")
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    with open(sys.argv[1], encoding="utf-8") as fh:
        spec = json.load(fh)
    build_xlsx(spec, sys.argv[2])
    for w in warnings:
        print("WARN", w)
    print(f"Wrote {sys.argv[2]}")

